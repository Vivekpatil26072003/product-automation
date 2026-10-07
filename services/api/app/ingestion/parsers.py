"""Native parsers (FR04). Output is plain data so it can cross the sandbox process boundary.

A page is the unit of evidence: a PDF page, a worksheet, or the whole of a TXT/DOCX file.
Every span records where it came from (character offsets always; sheet/cell or paragraph/table
coordinates for Office files). No bounding box is ever invented: native PDF text is page-level.

Nothing is executed: formulas are read from cached values only, links and macros are ignored.
"""

import io
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from app.ingestion import limits

NATIVE_PDF_MIN_CHARS = 20  # fewer non-space characters than this means the page needs OCR


class ParseError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


class _PageBuilder:
    """Accumulates page text and spans with exact character offsets."""

    def __init__(self, page_no: int, parser: str):
        self.page_no, self.parser = page_no, parser
        self.parts: list[str] = []
        self.length = 0
        self.spans: list[dict[str, Any]] = []
        self.warnings: list[str] = []

    def add(self, text: str, sep: str = "\n", **where: Any) -> None:
        if self.parts and sep:
            self.parts.append(sep)
            self.length += len(sep)
        start = self.length
        self.parts.append(text)
        self.length += len(text)
        span = {"id": f"p{self.page_no}-s{len(self.spans) + 1}", "page": self.page_no, "text": text,
                "char_start": start, "char_end": self.length, "polygon": None, "confidence": None}  # fmt: skip
        span.update(where)
        self.spans.append(span)

    def build(self) -> dict[str, Any]:
        return {"page_no": self.page_no, "parser": self.parser, "text": "".join(self.parts),
                "spans": self.spans, "warnings": sorted(set(self.warnings)), "needs_ocr": False}  # fmt: skip


def parse_txt(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
    page = _PageBuilder(1, "native-txt")
    # Keep true file offsets: lines are spans over the original text, blank lines included in offsets.
    offset = 0
    for line_no, line in enumerate(text.split("\n"), start=1):
        if line.strip():
            page.spans.append({"id": f"p1-s{len(page.spans) + 1}", "page": 1, "text": line,
                               "char_start": offset, "char_end": offset + len(line), "polygon": None,
                               "confidence": None, "line": line_no})  # fmt: skip
        offset += len(line) + 1
    out = page.build()
    out["text"] = text
    return [out]


def _cell_text(value: Any) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # openpyxl returns binary floats; repr() is the shortest exact round-trip text. Integers drop ".0".
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    pages, total_rows = [], 0
    for index, ws in enumerate(values.worksheets, start=1):
        fws = formulas.worksheets[index - 1]
        page = _PageBuilder(index, "native-xlsx")
        first_in_row = True
        for row, frow in zip(ws.iter_rows(), fws.iter_rows(), strict=False):
            if not any(c.value is not None for c in row) and not any(
                isinstance(c.value, str) and c.value.startswith("=") for c in frow
            ):
                continue
            total_rows += 1
            if total_rows > limits.MAX_OFFICE_ROWS:
                raise ParseError("ROW_LIMIT", f"The workbook has more than {limits.MAX_OFFICE_ROWS} rows.")
            first_in_row = True
            for cell, fcell in zip(row, frow, strict=False):
                is_formula = isinstance(fcell.value, str) and fcell.value.startswith("=")
                if cell.value is None:
                    if is_formula:
                        page.warnings.append("FORMULA_WITHOUT_CACHED_VALUE")
                        page.add("", sep="\n" if first_in_row and page.spans else "\t", sheet=ws.title,
                                 cell=cell.coordinate, issue="FORMULA_WITHOUT_CACHED_VALUE")  # fmt: skip
                        first_in_row = False
                    continue
                text = _cell_text(cell.value)
                if not text:
                    continue
                page.add(text, sep="\n" if first_in_row else "\t", sheet=ws.title, cell=cell.coordinate,
                         **({"formula_cached": True} if is_formula else {}))  # fmt: skip
                first_in_row = False
        pages.append(page.build())
    values.close()
    formulas.close()
    return pages


def parse_docx(data: bytes) -> list[dict[str, Any]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    page = _PageBuilder(1, "native-docx")
    for i, para in enumerate(document.paragraphs, start=1):
        if para.text.strip():
            page.add(para.text.strip(), paragraph=i)
    rows = 0
    for t_index, table in enumerate(document.tables, start=1):
        for r_index, row in enumerate(table.rows, start=1):
            rows += 1
            if rows > limits.MAX_OFFICE_ROWS:
                raise ParseError("ROW_LIMIT", f"The document has more than {limits.MAX_OFFICE_ROWS} table rows.")
            first = True
            for c_index, cell in enumerate(row.cells, start=1):
                text = cell.text.strip()
                if text:
                    page.add(text, sep="\n" if first else "\t", table=t_index, row=r_index, col=c_index)
                    first = False
    return [page.build()]


def parse_pdf(data: bytes) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    pages = []
    for n, pdf_page in enumerate(reader.pages, start=1):
        page = _PageBuilder(n, "native-pdf")
        try:
            text = pdf_page.extract_text() or ""
        except Exception:  # noqa: BLE001 - a broken text layer falls back to OCR
            text = ""
        for line in text.splitlines():
            if line.strip():
                page.add(line.strip())
        built = page.build()
        built["needs_ocr"] = sum(not c.isspace() for c in built["text"]) < NATIVE_PDF_MIN_CHARS
        pages.append(built)
    return pages


def render_pdf_page(data: bytes, page_no: int, dpi: int = 200) -> bytes:
    """Render one PDF page to PNG for OCR (derived copy; the original is never modified)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        bitmap = pdf[page_no - 1].render(scale=dpi / 72)
        image = bitmap.to_pil()
        if image.width * image.height > limits.MAX_IMAGE_PIXELS:
            image.thumbnail((6000, 6000))
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        return buf.getvalue()
    finally:
        pdf.close()


def normalize_image(data: bytes) -> bytes:
    """Derived processing image: EXIF orientation applied, all metadata (incl. GPS) dropped (spec §8)."""
    from PIL import Image, ImageOps

    with Image.open(io.BytesIO(data)) as im:
        upright = ImageOps.exif_transpose(im)
        clean = Image.new(upright.mode if upright.mode in ("RGB", "L") else "RGB", upright.size)
        clean.paste(upright.convert(clean.mode))
        buf = io.BytesIO()
        clean.save(buf, format="PNG")
        return buf.getvalue()


NATIVE = {"txt": parse_txt, "xlsx": parse_xlsx, "docx": parse_docx, "pdf": parse_pdf}
