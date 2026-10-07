"""Typed XLSX rendering from an export snapshot (FR14, spec §11 "Excel files", TC29).

- Records, Summary and Metadata sheets.
- Every text cell is written explicitly as a string, so note text such as "=HYPERLINK(...)", "+SUM" or
  "@cmd" stays inert text and never becomes a formula. No cell is ever written as a formula.
- Quantities are numeric when the decimal value round-trips exactly through Excel's double precision;
  otherwise they are written as documented decimal text (never silently rounded).
- Dates are real dates with an explicit yyyy-mm-dd format.
The workbook is a downloaded copy: it never synchronizes back into the system.
"""

import io
import json
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATE_FORMAT = "yyyy-mm-dd"
QTY_FORMAT = "#,##0.000"
PCT_FORMAT = "0.0"


def _text(ws, row: int, col: int, value: str | None) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"  # force a string cell even when the text starts with =, +, - or @


def _qty(ws, row: int, col: int, value: str | None) -> None:
    if value is None:
        return _text(ws, row, col, "")
    d = Decimal(value)
    if Decimal(repr(float(d))) == d:
        cell = ws.cell(row=row, column=col, value=float(d))
        cell.number_format = QTY_FORMAT
    else:
        _text(ws, row, col, format(d, "f"))  # beyond double precision: exact decimal text


def _date(ws, row: int, col: int, value: str) -> None:
    cell = ws.cell(row=row, column=col, value=date.fromisoformat(value))
    cell.number_format = DATE_FORMAT


def _headers(ws, names: list[str]) -> None:
    for i, name in enumerate(names, start=1):
        _text(ws, 1, i, name)
        ws.cell(row=1, column=i).font = Font(bold=True)
    ws.freeze_panes = "A2"


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


RECORD_COLUMNS = [
    "Date",
    "Department",
    "Machine",
    "Operator",
    "Production",
    "Target",
    "Unit",
    "Achievement %",
    "Status",
    "Stop minutes",
    "Remarks",
    "State",
    "Record ID",
    "Revision",
]


def render(export: dict[str, Any]) -> bytes:
    """`export` = {id, filter, rows, metrics, data_version, timezone, created_at, row_count}."""
    wb = Workbook()
    wb.properties.creator = "Production Automation"
    wb.properties.title = "Approved production records"

    ws = wb.active
    ws.title = "Records"
    _headers(ws, RECORD_COLUMNS)
    for i, rec in enumerate(export["rows"], start=2):
        _date(ws, i, 1, rec["production_date"])
        _text(ws, i, 2, rec["department_name"])
        _text(ws, i, 3, rec["machine_code"])
        _text(ws, i, 4, rec["operator_name"])
        _qty(ws, i, 5, rec["production_qty"])
        _qty(ws, i, 6, rec["target_qty"])
        _text(ws, i, 7, rec["unit"])
        if rec["achievement_pct"] is None:
            _text(ws, i, 8, "N/A")
        else:
            ws.cell(row=i, column=8, value=float(rec["achievement_pct"])).number_format = PCT_FORMAT
        _text(ws, i, 9, rec["status"])
        ws.cell(row=i, column=10, value=int(rec["stop_minutes"]))
        _text(ws, i, 11, rec["remarks"])
        _text(ws, i, 12, rec["state"])
        _text(ws, i, 13, rec["record_id"])
        ws.cell(row=i, column=14, value=int(rec["revision"]))
    _widths(ws, [12, 16, 10, 18, 14, 14, 7, 13, 12, 12, 40, 10, 38, 9])

    ws = wb.create_sheet("Summary")
    metrics = export["metrics"]
    _headers(ws, ["Scope", "Unit", "Production", "Target", "Achievement %", "Variance", "Records"])
    row = 2
    for scope, items in (
        ("All selected", metrics["metrics"]),
        *[(dep["department_name"], [dep]) for dep in metrics["departments"]],
    ):
        for x in items:
            _text(ws, row, 1, scope)
            _text(ws, row, 2, x["unit"])
            _qty(ws, row, 3, x["production_qty"])
            _qty(ws, row, 4, x["target_qty"])
            if x["achievement_pct"] is None:
                _text(ws, row, 5, "N/A")
            else:
                ws.cell(row=row, column=5, value=float(x["achievement_pct"])).number_format = PCT_FORMAT
            _qty(ws, row, 6, x["variance"])
            ws.cell(row=row, column=7, value=x["record_count"])
            row += 1
    row += 1
    _text(ws, row, 1, "Status")
    _text(ws, row, 2, "Records")
    for status, n in metrics["status_counts"].items():
        row += 1
        _text(ws, row, 1, status)
        ws.cell(row=row, column=2, value=n)
    row += 2
    _text(ws, row, 1, "Record downtime minutes (may overlap across machines; not plant downtime)")
    ws.cell(row=row, column=2, value=metrics["stop_total_minutes"])
    _widths(ws, [28, 7, 14, 14, 13, 14, 9])

    ws = wb.create_sheet("Metadata")
    meta = [
        ("Export ID", export["id"]),
        ("Generated (UTC)", export["created_at"]),
        ("Company time zone", export["timezone"]),
        ("Data version", str(export["data_version"])),
        ("Records", str(export["row_count"])),
        ("Filter", json.dumps(export["filter"], sort_keys=True)),
        (
            "Contents",
            "Approved production records at their current revision when the export was requested. "
            "Archived records are included only if the filter says so. Quantities in m, kg and pcs are "
            "never added together. This file is a copy: edits here do not change the system.",
        ),
    ]
    for i, (k, v) in enumerate(meta, start=1):
        _text(ws, i, 1, k)
        ws.cell(row=i, column=1).font = Font(bold=True)
        _text(ws, i, 2, v)
    _widths(ws, [20, 100])
    ws.protection.sheet = True  # guard against accidental edits of the provenance

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def created_at_text(value: datetime) -> str:
    return value.isoformat(timespec="seconds")
