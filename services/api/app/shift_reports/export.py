"""Files of one daily sheet: Excel (with formulas, like the company's sheet), PDF and CSV.

Excel: sheet "Report" has one block per section; the Total column and the calculated rows are Excel formulas that
reference the shift cells and the table parameters written above each loom table, so the workbook keeps working
if someone corrects a value in Excel. To date is written as a value (it needs the other days of the month).
Sheet "Data" lists every saved value as one row (the database rows), sheet "Notes" the notes and supervisors.
All text is written as text cells: nothing written on a page can become a formula.
"""

import csv
import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import Connection, select

from app.db import tables as t
from app.orders.pdf import _STYLES, BOLD, _Text
from app.reports.pdf import HEAD_BG, RULE
from app.shift_reports import catalog, compute

MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "csv": "text/csv",
}
HEAD = PatternFill("solid", fgColor="EEF2F7")
CALC = PatternFill("solid", fgColor="F7F9FC")


def _num(v: Decimal | None) -> float | None:
    return None if v is None else float(v)


def _context(conn: Connection, row: Any) -> dict[str, Any]:
    from app.owner_reports.settings import company_name
    from app.shift_reports.service import grids_for

    grids, todate, raw = grids_for(conn, row)
    targets, params = raw.pop("_targets"), raw.pop("_params")
    dept = conn.execute(select(t.department.c.name).where(t.department.c.id == row.department_id)).scalar_one()
    return {
        "grids": grids,
        "todate": todate,
        "raw": raw,
        "targets": targets,
        "params": params,
        "dept": dept,
        "company": company_name(conn, row.tenant_id),
    }


def file_name(row: Any, dept: str, fmt: str) -> str:
    safe = "".join(ch if ch.isalnum() else "_" for ch in dept)[:30]
    return f"Daily_Sheet_{row.report_date:%Y-%m-%d}_{safe}.{fmt}"


def render(conn: Connection, row: Any, fmt: str) -> tuple[bytes, str, str]:
    c = _context(conn, row)
    data = {"xlsx": _xlsx, "pdf": _pdf, "csv": _csv}[fmt](row, c)
    return data, file_name(row, c["dept"], fmt), MIME[fmt]


# --- Excel -----------------------------------------------------------------------------------


def _text(ws, r: int, col: int, value: Any, bold: bool = False) -> None:
    cell = ws.cell(row=r, column=col)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    if bold:
        cell.font = Font(bold=True)


def _xlsx(row: Any, c: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    _text(ws, 1, 1, f"{c['company']} - daily production sheet", True)
    _text(
        ws,
        2,
        1,
        f"Date: {row.report_date:%d %b %Y}  ·  Department: {c['dept']}  ·  "
        f"{'Approved' if row.state == 'APPROVED' else 'NOT YET APPROVED (draft)'}",
    )
    sups = "  ·  ".join(
        f"Shift {sh}: {(row.shifts or {}).get(sh, {}).get('supervisor') or '-'}" for sh in catalog.SHIFTS
    )
    _text(ws, 3, 1, sups)
    r = 5
    for s in catalog.SECTIONS:
        day = s.shifts == ("D",)
        _text(ws, r, 1, s.title, True)
        r += 1
        prefs: dict[str, str] = {}
        if s.params:  # table parameters used by the formulas below
            col = 2
            for p, v in c["params"][s.key].items():
                _text(ws, r, col, catalog.PARAM_LABEL[p])
                ws.cell(row=r, column=col + 1, value=float(v))
                prefs[p] = f"${get_column_letter(col + 1)}${r}"
                col += 2
            r += 1
        heads = (
            ["Row", "Unit", "Target", "Today", "To date"]
            if day
            else ["Row", "Unit", "Target", "I", "II", "III", "Total", "To date"]
        )
        for i, h in enumerate(heads, start=1):
            _text(ws, r, i, h, True)
            ws.cell(row=r, column=i).fill = HEAD
        r += 1
        first = r
        rows_at: dict[str, int] = {}
        for m in s.metrics:
            rows_at[m.key] = r
            r += 1
        for m in s.metrics:
            rr = rows_at[m.key]
            _text(ws, rr, 1, m.label)
            _text(ws, rr, 2, m.unit)
            target = c["targets"][(s.key, m.key)]
            if target is not None:
                ws.cell(row=rr, column=3, value=float(target))
            cols = ["D"] if day else ["D", "E", "F"]
            for col, sh in zip(cols, s.shifts, strict=True):
                cell = ws[f"{col}{rr}"]
                if m.kind == "input":
                    cell.value = _num(c["grids"][s.key][m.key][sh])
                else:
                    cell.value = _formula(s, m.key, col, rows_at, first, rr, prefs)
                    cell.fill = CALC
            if day:
                todate_col = "E"
            else:
                fn = "AVERAGE" if m.agg == "avg" else "SUM"
                ws[f"G{rr}"] = f'=IF(COUNT(D{rr}:F{rr})=0,"",{fn}(D{rr}:F{rr}))'
                ws[f"G{rr}"].fill = CALC
                todate_col = "H"
            ws[f"{todate_col}{rr}"] = _num(c["todate"][(s.key, m.key)])
        r += 1
    for col, width in zip("ABCDEFGH", (38, 6, 12, 12, 12, 12, 13, 13), strict=True):
        ws.column_dimensions[col].width = width
    for row_cells in ws.iter_rows(min_row=5, min_col=3, max_col=8):
        for cell in row_cells:
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")

    data = wb.create_sheet("Data")  # one row per saved value: what is stored in the database
    for i, h in enumerate(["Date", "Department", "Section", "Row", "Shift", "Supervisor", "Value", "Entered by"], 1):
        _text(data, 1, i, h, True)
    n = 2
    for (sec, met, sh), v in sorted(c["raw"].items()):
        s = catalog.BY_KEY[sec]
        data.cell(row=n, column=1, value=row.report_date)
        data.cell(row=n, column=1).number_format = "yyyy-mm-dd"
        for i, val in enumerate(
            [c["dept"], s.title, s.metric(met).label, sh, (row.shifts or {}).get(sh, {}).get("supervisor") or ""],
            start=2,
        ):
            _text(data, n, i, val)
        data.cell(row=n, column=7, value=_num(v.value))
        _text(
            data,
            n,
            8,
            {"read": "read from page", "ai": "read by AI", "reviewer": "reviewer", "manual": "entered"}.get(
                v.source, v.source
            ),
        )
        n += 1
    notes = wb.create_sheet("Notes")
    _text(notes, 1, 1, "Label", True)
    _text(notes, 1, 2, "Text", True)
    for i, note in enumerate(row.notes or [], start=2):
        _text(notes, i, 1, note.get("label"))
        _text(notes, i, 2, note.get("text"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _formula(
    s: catalog.Section, key: str, col: str, rows_at: dict[str, int], first: int, rr: int, p: dict[str, str]
) -> str | None:
    if key == "total":
        top, bottom = first, rr - 1
        return f'=IF(COUNT({col}{top}:{col}{bottom})=0,"",SUM({col}{top}:{col}{bottom}))'
    if key == "meters_per_min":
        m, h = f"{col}{rows_at['meters']}", f"{col}{rows_at['working_hours']}"
        return f'=IF(OR({m}="",{h}="",{h}=0),"",{m}/({h}*60))'
    run, picks = f"{col}{rows_at['running_looms']}", f"{col}{rows_at['picks']}"
    formulas = {
        "theoretical_picks": f'=IF({run}="","",{run}*{p.get("theo_rate")}*8)',
        "loss_of_pick": f'=IF(OR({run}="",{picks}=""),"",{run}*{p.get("theo_rate")}*8-{picks})',
        "utilization_pct": f'=IF({run}="","",{run}/{p.get("installed")}*100)',
        "working_pct": f'=IF(OR({run}="",{picks}="",{run}=0),"",{picks}/({run}*{p.get("rate")}*8)*100)',
        "total_eff_pct": f'=IF({picks}="","",{picks}/({p.get("installed")}*{p.get("rate")}*8)*100)',
        "picks_per_hour": f'=IF(OR({run}="",{picks}="",{run}=0),"",{picks}/({run}*8))',
    }
    return formulas.get(key)


# --- CSV -------------------------------------------------------------------------------------


def _csv(row: Any, c: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(
        [
            "date",
            "department",
            "section",
            "row",
            "unit",
            "target",
            "shift_I",
            "shift_II",
            "shift_III",
            "day",
            "total",
            "to_date",
            "calculated",
        ]
    )

    def cell(v: Decimal | None) -> str:
        text = "" if v is None else format(v.normalize(), "f")
        return text

    for s in catalog.SECTIONS:
        for m in s.metrics:
            g = c["grids"][s.key][m.key]
            label = m.label
            if label[:1] in "=+-@":  # never a formula when opened in a spreadsheet
                label = "'" + label
            w.writerow(
                [
                    row.report_date.isoformat(),
                    c["dept"],
                    s.title,
                    label,
                    m.unit,
                    cell(c["targets"][(s.key, m.key)]),
                    *(cell(g.get(sh)) for sh in catalog.SHIFTS),
                    cell(g.get("D")),
                    cell(g["total"]),
                    cell(c["todate"][(s.key, m.key)]),
                    "yes" if m.kind == "derived" else "no",
                ]
            )
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens Gujarati / Hindi text correctly


# --- PDF -------------------------------------------------------------------------------------


def _pdf(row: Any, c: dict[str, Any]) -> bytes:
    tx, sty = _Text(), _STYLES
    P = lambda text, style="cell": Paragraph(tx(text), sty[style])  # noqa: E731
    sups = " · ".join(f"Shift {sh}: {(row.shifts or {}).get(sh, {}).get('supervisor') or '-'}" for sh in catalog.SHIFTS)
    story: list[Any] = [
        P(f"Daily production sheet - {c['company']}", "title"),
        P(
            f"{row.report_date:%d %b %Y} · {c['dept']} · "
            f"{'approved' if row.state == 'APPROVED' else 'NOT YET APPROVED (draft)'} · {sups}",
            "sub",
        ),
        Spacer(1, 3 * mm),
    ]
    for s in catalog.SECTIONS:
        day = s.shifts == ("D",)
        heads = (
            ["Row", "Target", "Today", "To date"] if day else ["Row", "Target", "I", "II", "III", "Total", "To date"]
        )
        rows = [[P(h, "key") for h in heads]]
        any_value = False
        for m in s.metrics:
            g = c["grids"][s.key][m.key]
            shift_vals = [g.get(sh) for sh in s.shifts]
            any_value = any_value or any(v is not None for v in shift_vals)
            cells = [compute.rounded(v) for v in shift_vals]
            line = [
                P(m.label + (" *" if m.kind == "derived" else "")),
                P(compute.rounded(c["targets"][(s.key, m.key)])),
                *(P(x) for x in cells),
            ]
            if not day:
                line.append(P(compute.rounded(g["total"])))
            line.append(P(compute.rounded(c["todate"][(s.key, m.key)])))
            rows.append(line)
        if not any_value:
            continue  # sections with nothing written are left out of the printout
        widths = [80 * mm, 26 * mm, 30 * mm, 30 * mm] if day else [78 * mm] + [26 * mm] * 6
        table = Table(rows, colWidths=widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), HEAD_BG),
                    ("FONTNAME", (0, 0), (-1, 0), BOLD),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, RULE),
                    ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(KeepTogether([P(s.title, "key"), table, Spacer(1, 3 * mm)]))
    if row.notes:
        story.append(P("Notes", "key"))
        for n in row.notes:
            story.append(P(f"{n.get('label')}: {n.get('text')}"))
    story.append(Spacer(1, 3 * mm))
    story.append(P("* calculated from the written values. To date = average of the daily totals this month.", "sub"))
    if tx.replaced:
        story.append(P(f"{tx.replaced} character(s) could not be printed in this font and appear as ?.", "sub"))
    buf = io.BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=f"Daily sheet {row.report_date}",
        author=c["company"],
        invariant=1,
        pageCompression=1,
    ).build(story)
    return buf.getvalue()
