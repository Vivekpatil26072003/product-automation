"""Files of one pick reading register: Excel, PDF, CSV and SQL.

Excel: one sheet per shift laid out like the register page (machine rows; reading and picks per time); the machine
Total column and the calculated column totals are Excel formulas, so the workbook keeps working if someone corrects
a value in Excel. Below them: the totals the worker wrote and the stopped machines. Sheet "Data" lists every saved
value as one row (the database rows), sheet "Checks" the open checks and notes.
SQL: CREATE TABLE IF NOT EXISTS + DELETE of this day + INSERT rows, in standard SQL that PostgreSQL, MySQL / MariaDB
and SQLite load as they are. Text from the page is escaped (quotes doubled, backslashes and control characters
removed), so nothing written on a page can change the statements.
All spreadsheet text is written as text cells: nothing written on a page can become a formula.
"""

import csv
import io
import re
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import Connection

from app.orders.pdf import _STYLES, BOLD, _Text
from app.pick_registers import compute
from app.pick_registers.layout import FORM, SHIFT_HOURS, SHIFTS, SLOTS, TIMES, machine_sort, status_label
from app.reports.pdf import HEAD_BG, RULE

MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "csv": "text/csv",
    "sql": "application/sql",
}
HEAD = PatternFill("solid", fgColor="EEF2F7")
CALC = PatternFill("solid", fgColor="F7F9FC")
CHECK = PatternFill("solid", fgColor="FFF4D6")
SOURCE = {"read": "read from page", "ai": "read by AI", "reviewer": "reviewer", "manual": "entered"}


def _num(v: Decimal | None) -> float | int | None:
    if v is None:
        return None
    return int(v) if v == v.to_integral() else float(v)


def _s(v: Decimal | None) -> str:
    return "" if v is None else format(v.normalize(), "f")


def file_name(row: Any, dept: str, fmt: str) -> str:
    safe = "".join(ch if ch.isalnum() else "_" for ch in dept)[:30]
    return f"Pick_Register_{row.register_date:%Y-%m-%d}_{safe}.{fmt}"


def render(conn: Connection, row: Any, fmt: str) -> tuple[bytes, str, str]:
    from app.pick_registers.service import export_context

    c = export_context(conn, row)
    data = {"xlsx": _xlsx, "pdf": _pdf, "csv": _csv, "sql": _sql}[fmt](row, c)
    return data, file_name(row, c["dept"], fmt), MIME[fmt]


def _shifts(c: dict[str, Any]) -> list[str]:
    return [sh for sh in SHIFTS if c["res"].machines[sh] or any(k[0] == sh for k in c["written"])]


def _open_checks(c: dict[str, Any]) -> list[tuple[str, str, int, str]]:
    out = []
    for key in sorted(c["res"].issues, key=lambda k: (SHIFTS.index(k[0]), machine_sort(k[1]), k[2])):
        v = c["values"].get(key)
        for i in compute.open_issues(c["res"].issues[key], v.accepted if v is not None else None):
            out.append((*key, i.text))
    for key, v in sorted(c["values"].items(), key=lambda x: (SHIFTS.index(x[0][0]), machine_sort(x[0][1]), x[0][2])):
        if v.uncertain:
            out.append((*key, v.note or "Check against the photo."))
    for (sh, k), issue in sorted(c["res"].total_issues.items(), key=lambda x: (SHIFTS.index(x[0][0]), x[0][1])):
        w = c["written"].get((sh, k))
        if compute.open_issues([issue], w.accepted if w is not None else None):
            out.append((sh, "TOTAL", k, issue.text))
    return out


def _where(sh: str, m: str, k: int) -> str:
    return f"Shift {sh}, total {TIMES[sh][k]}" if m == "TOTAL" else f"Shift {sh}, m/c {m}, {TIMES[sh][k]}"


# --- Excel -----------------------------------------------------------------------------------


def _text(ws, r: int, col: int, value: Any, bold: bool = False) -> None:
    cell = ws.cell(row=r, column=col)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    if bold:
        cell.font = Font(bold=True)


def _xlsx(row: Any, c: dict[str, Any]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    res, values, written = c["res"], c["values"], c["written"]
    flagged = {(sh, m, k) for sh, m, k, _ in _open_checks(c)}
    state = "Approved" if row.state == "APPROVED" else "NOT YET APPROVED (draft)"
    for sh in _shifts(c) or ["II"]:
        ws = wb.create_sheet(f"Shift {sh}")
        _text(ws, 1, 1, f"{c['company']} - Hourly production reading register ({FORM}) - pick reading", True)
        _text(ws, 2, 1, f"Date: {row.register_date:%d %b %Y}  ·  Department: {c['dept']}  ·  Shift {sh} "
                        f"({SHIFT_HOURS[sh]})  ·  {state}")  # fmt: skip
        # Columns: A machine | B start reading | C,D reading/picks | E,F | G,H | I,J | K total | L marks
        heads = ["M/c No.", f"{TIMES[sh][0]} reading"]
        for k in SLOTS[1:]:
            heads += [f"{TIMES[sh][k]} reading", f"{TIMES[sh][k]} picks"]
        heads += ["Total picks", "Marks"]
        for i, h in enumerate(heads, start=1):
            _text(ws, 4, i, h, True)
            ws.cell(row=4, column=i).fill = HEAD
        r = 5
        first = r
        picks_cols = [get_column_letter(3 + 2 * (k - 1) + 1) for k in SLOTS[1:]]  # D F H J
        for m in res.machines[sh]:
            _text(ws, r, 1, m)
            marks = []
            for k in SLOTS:
                v = values.get((sh, m, k))
                if v is None:
                    continue
                rc = 2 if k == 0 else 3 + 2 * (k - 1)
                if v.reading is not None:
                    ws.cell(row=r, column=rc, value=_num(v.reading))
                if k and v.picks is not None:
                    ws.cell(row=r, column=rc + 1, value=_num(v.picks))
                if v.status:
                    marks.append(f"{TIMES[sh][k]} {status_label(v.status)}")
                if (sh, m, k) in flagged:
                    for col in (rc, rc + 1) if k else (rc,):
                        ws.cell(row=r, column=col).fill = CHECK
            ws[f"K{r}"] = f'=IF(COUNT({",".join(f"{p}{r}" for p in picks_cols)})=0,"",' \
                          f'SUM({",".join(f"{p}{r}" for p in picks_cols)}))'  # fmt: skip
            ws[f"K{r}"].fill = CALC
            _text(ws, r, 12, "; ".join(marks))
            r += 1
        last = r - 1
        _text(ws, r, 1, "Total (calculated)", True)
        for p in picks_cols + ["K"]:
            ws[f"{p}{r}"] = f'=IF(COUNT({p}{first}:{p}{last})=0,"",SUM({p}{first}:{p}{last}))' if last >= first else ""
            ws[f"{p}{r}"].fill = CALC
        r += 1
        _text(ws, r, 1, "Total (written)", True)
        for k in SLOTS:
            w = written.get((sh, k))
            if w is not None and w.written is not None:
                col = "B" if k == 0 else picks_cols[k - 1]
                ws[f"{col}{r}"] = _num(w.written)
        if (sh, 0) in written:
            _text(ws, r, 12, f"Under {TIMES[sh][0]}: shift / day total as written")
        r += 1
        _text(ws, r, 1, "M/c stop", True)
        for k in SLOTS:
            col = "B" if k == 0 else picks_cols[k - 1]
            ws[f"{col}{r}"] = res.stopped[(sh, k)]
        ws.column_dimensions["A"].width = 18
        for i in range(2, 12):
            ws.column_dimensions[get_column_letter(i)].width = 11
        ws.column_dimensions["L"].width = 30
        for row_cells in ws.iter_rows(min_row=5, min_col=2, max_col=11):
            for cell in row_cells:
                cell.number_format = "#,##0.####"
                cell.alignment = Alignment(horizontal="right")
        ws.freeze_panes = "B5"

    data = wb.create_sheet("Data")  # one row per saved value: what is stored in the database
    for i, h in enumerate(
        ["Date", "Department", "Shift", "Time", "Machine", "Reading", "Picks", "Mark", "Entered by"], 1
    ):
        _text(data, 1, i, h, True)
    n = 2
    for (sh, m, k), v in sorted(values.items(), key=lambda x: (SHIFTS.index(x[0][0]), machine_sort(x[0][1]), x[0][2])):
        data.cell(row=n, column=1, value=row.register_date).number_format = "yyyy-mm-dd"
        for i, val in enumerate([c["dept"], sh, TIMES[sh][k], m], start=2):
            _text(data, n, i, val)
        data.cell(row=n, column=6, value=_num(v.reading))
        data.cell(row=n, column=7, value=_num(v.picks))
        _text(data, n, 8, status_label(v.status))
        _text(data, n, 9, SOURCE.get(v.source, v.source))
        n += 1
    checks = wb.create_sheet("Checks")
    _text(checks, 1, 1, "Where", True)
    _text(checks, 1, 2, "Check", True)
    n = 2
    for sh, m, k, text in _open_checks(c):
        _text(checks, n, 1, _where(sh, m, k))
        _text(checks, n, 2, text)
        n += 1
    for note in res.notes + (row.notes or []):
        _text(checks, n, 1, note.get("label"))
        _text(checks, n, 2, note.get("text"))
        n += 1
    checks.column_dimensions["A"].width = 28
    checks.column_dimensions["B"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- CSV -------------------------------------------------------------------------------------


def _safe(text: str) -> str:
    return "'" + text if text[:1] and text[:1] in "=+-@" else text  # never a formula when opened in a spreadsheet


def _csv(row: Any, c: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["date", "department", "shift", "time", "slot", "machine", "record", "reading", "picks", "mark",
                "entered_by"])  # fmt: skip
    res, values = c["res"], c["values"]
    for (sh, m, k), v in sorted(values.items(), key=lambda x: (SHIFTS.index(x[0][0]), machine_sort(x[0][1]), x[0][2])):
        w.writerow([row.register_date.isoformat(), _safe(c["dept"]), sh, TIMES[sh][k], k, _safe(m), "value",
                    _s(v.reading), _s(v.picks), _safe(status_label(v.status)), v.source])  # fmt: skip
    for sh in _shifts(c):
        for k in SLOTS:
            head = [row.register_date.isoformat(), _safe(c["dept"]), sh, TIMES[sh][k]]
            if k:
                calc = _s(res.column_total[(sh, k)])
                w.writerow([*head, k, "", "total_calculated", "", calc, "", "calculated"])
            wt = c["written"].get((sh, k))
            if wt is not None:
                kind = "total_written" if k else "shift_total_written"
                w.writerow([*head, k, "", kind, "", _s(wt.written), "", wt.source])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens Gujarati / Hindi text correctly


# --- SQL -------------------------------------------------------------------------------------


def _q(text: Any) -> str:
    if text is None or text == "":
        return "NULL"
    clean = re.sub(r"[\x00-\x1f\x7f]", " ", str(text)).replace("\\", "/").replace("'", "''")
    return f"'{clean}'"


def _n(v: Decimal | None) -> str:
    return "NULL" if v is None else format(v.normalize(), "f")


def _sql(row: Any, c: dict[str, Any]) -> bytes:
    res, values, day, dept = c["res"], c["values"], row.register_date.isoformat(), c["dept"]
    lines = [
        f"-- {c['company']}: hourly production reading register ({FORM}), pick reading",
        f"-- Date {day}, department {re.sub(r'[\x00-\x1f]', ' ', dept)}, "
        f"{'approved' if row.state == 'APPROVED' else 'NOT YET APPROVED (draft)'}, version {row.version}",
        f"-- Written {datetime.now(UTC):%Y-%m-%d %H:%M} UTC. Values as written on the register; totals calculated.",
        "-- Standard SQL: loads into PostgreSQL, MySQL / MariaDB and SQLite. Running it again replaces this day.",
        "",
        "CREATE TABLE IF NOT EXISTS pick_reading (",
        "  register_date DATE NOT NULL,",
        "  department VARCHAR(120) NOT NULL,",
        "  shift VARCHAR(3) NOT NULL,",
        "  reading_time VARCHAR(5) NOT NULL,",
        "  slot SMALLINT NOT NULL,",
        "  machine VARCHAR(20) NOT NULL,",
        "  meter_reading DECIMAL(18,4),",
        "  picks DECIMAL(18,4),",
        "  mark VARCHAR(40),",
        "  entered_by VARCHAR(20) NOT NULL,",
        "  PRIMARY KEY (register_date, department, shift, machine, slot)",
        ");",
        "CREATE TABLE IF NOT EXISTS pick_reading_total (",
        "  register_date DATE NOT NULL,",
        "  department VARCHAR(120) NOT NULL,",
        "  shift VARCHAR(3) NOT NULL,",
        "  reading_time VARCHAR(5) NOT NULL,",
        "  slot SMALLINT NOT NULL,",
        "  calculated_picks DECIMAL(18,4),",
        "  written_total DECIMAL(18,4),",
        "  stopped_machines INTEGER NOT NULL,",
        "  PRIMARY KEY (register_date, department, shift, slot)",
        ");",
        "",
        # Text of the downloaded file, not a query run here; the date is ISO and the department is escaped.
        f"DELETE FROM pick_reading WHERE register_date = '{day}' AND department = {_q(dept)};",  # noqa: S608
        f"DELETE FROM pick_reading_total WHERE register_date = '{day}' AND department = {_q(dept)};",  # noqa: S608
    ]
    rows = [
        f"('{day}', {_q(dept)}, '{sh}', '{TIMES[sh][k]}', {k}, {_q(m)}, {_n(v.reading)}, {_n(v.picks)}, "
        f"{_q(status_label(v.status))}, {_q(v.source)})"
        for (sh, m, k), v in sorted(
            values.items(), key=lambda x: (SHIFTS.index(x[0][0]), machine_sort(x[0][1]), x[0][2])
        )
    ]
    if rows:
        lines.append(
            "INSERT INTO pick_reading (register_date, department, shift, reading_time, slot, machine, meter_reading, "
            "picks, mark, entered_by) VALUES"
        )
        lines.append(",\n".join(rows) + ";")
    totals = []
    for sh in _shifts(c):
        for k in SLOTS:
            wt = c["written"].get((sh, k))
            calc = res.column_total[(sh, k)] if k else res.shift_total[sh]
            totals.append(
                f"('{day}', {_q(dept)}, '{sh}', '{TIMES[sh][k]}', {k}, {_n(calc)}, "
                f"{_n(wt.written if wt is not None else None)}, {res.stopped[(sh, k)]})"
            )
    if totals:
        lines.append(
            "INSERT INTO pick_reading_total (register_date, department, shift, reading_time, slot, calculated_picks, "
            "written_total, stopped_machines) VALUES"
        )
        lines.append(",\n".join(totals) + ";")
    lines.append("-- slot 0 of pick_reading_total: calculated_picks = shift total; written_total = the figure written")
    lines.append("-- under the first column (shift / day total).")
    return ("\n".join(lines) + "\n").encode("utf-8")


# --- PDF -------------------------------------------------------------------------------------


def _pdf(row: Any, c: dict[str, Any]) -> bytes:
    tx = _Text()
    sty = _STYLES | {  # 30+ machine rows fit one landscape page
        "cell": ParagraphStyle("r-cell", parent=_STYLES["cell"], fontSize=8, leading=9.6),
        "key": ParagraphStyle("r-key", parent=_STYLES["key"], fontSize=8, leading=9.6),
    }
    P = lambda text, style="cell": Paragraph(tx(text), sty[style])  # noqa: E731
    res, values, written = c["res"], c["values"], c["written"]
    flagged = {(sh, m, k) for sh, m, k, _ in _open_checks(c)}
    state = "approved" if row.state == "APPROVED" else "NOT YET APPROVED (draft)"
    story: list[Any] = []
    for n, sh in enumerate(_shifts(c)):
        if n:
            story.append(PageBreak())
        story += [
            P(f"Hourly production reading register ({FORM}) - pick reading - {c['company']}", "title"),
            P(f"{row.register_date:%d %b %Y} · {c['dept']} · shift {sh} ({SHIFT_HOURS[sh]}) · {state}", "sub"),
            Spacer(1, 2 * mm),
        ]
        rows = [[P(h, "key") for h in ["M/c", *TIMES[sh], "Total"]]]
        styles = []
        for i, m in enumerate(res.machines[sh], start=1):
            line = [P(m)]
            for k in SLOTS:
                v = values.get((sh, m, k))
                text = ""
                if v is not None:
                    parts = [_s(v.reading)]
                    if k and v.picks is not None:
                        parts.append(f"({_s(v.picks)})")
                    if v.status:
                        parts.append(status_label(v.status))
                    text = " ".join(p for p in parts if p)
                line.append(P(text))
                if (sh, m, k) in flagged:
                    styles.append(("BACKGROUND", (k + 1, i), (k + 1, i), "#FFF4D6"))
            line.append(P(_s(res.machine_total[(sh, m)])))
            rows.append(line)
        calc = [P("Total (calculated)", "key"), P("")]
        calc += [P(_s(res.column_total[(sh, k)])) for k in SLOTS[1:]] + [P(_s(res.shift_total[sh]), "key")]
        wrote = [P("Total (written)", "key")]
        wrote += [P(_s(written[(sh, k)].written) if (sh, k) in written else "") for k in SLOTS] + [P("")]
        stop = [P("M/c stop", "key")] + [P(str(res.stopped[(sh, k)])) for k in SLOTS] + [P("")]
        rows += [calc, wrote, stop]
        table = Table(rows, colWidths=[24 * mm] + [38 * mm] * 5 + [30 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), HEAD_BG),
                    ("FONTNAME", (0, 0), (-1, 0), BOLD),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.3, RULE),
                    ("LINEABOVE", (0, -3), (-1, -3), 0.8, RULE),
                    ("TOPPADDING", (0, 0), (-1, -1), 0.6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0.6),
                    *styles,
                ]
            )
        )
        story.append(table)
    if not story:
        story.append(P("Nothing is entered in this register yet.", "title"))
    story.append(Spacer(1, 3 * mm))
    story.append(P(f"Day total picks: {_s(res.day_total) or '-'}. Cells read as \"reading (picks)\"; "
                   "highlighted cells still need a person's check.", "sub"))  # fmt: skip
    checks = _open_checks(c)
    if checks or res.notes or row.notes:
        story.append(P("Checks and notes", "key"))
        for sh, m, k, text in checks[:60]:
            story.append(P(f"{_where(sh, m, k)}: {text}"))
        for note in res.notes + (row.notes or []):
            story.append(P(f"{note.get('label')}: {note.get('text')}"))
    if tx.replaced:
        story.append(P(f"{tx.replaced} character(s) could not be printed in this font and appear as ?.", "sub"))
    buf = io.BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
        title=f"Pick register {row.register_date}",
        author=c["company"],
        invariant=1,
        pageCompression=1,
    ).build(story)
    return buf.getvalue()
