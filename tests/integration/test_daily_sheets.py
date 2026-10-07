"""Daily production sheet end to end: notebook photo -> OCR lines -> values in the day's sheet (database) ->
review (correct / confirm) -> approve -> sheet list -> Excel / PDF / CSV -> email with the file attached.
OCR is a stand-in returning lines in the Azure adapter's shape; EmailJS is the recorded transport of test_orders.
"""

import base64
import csv
import io
import uuid

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.db import tables as t
from tests import filegen
from tests.conftest import idem, sign_in
from tests.integration.test_ingestion_pipeline import complete_all, start_batch
from tests.integration.test_orders import FakeEmailJs, NoteOcr, _infra, configure_email  # noqa: F401
from tests.unit.test_daily_sheet import NOTEBOOK
from workers.runtime import run_one

pytestmark = [pytest.mark.db, pytest.mark.infra]
WORK = ["upload.scan", "upload.parse", "upload.extract", "sheet.email"]


def drain() -> None:
    while run_one(WORK, "test-worker"):
        pass


def photo(client, seeded, monkeypatch, lines, name="diary-page.png", confidence=0.97, width=420) -> dict:
    monkeypatch.setattr("workers.ingestion.get_ocr", lambda: NoteOcr(lines, confidence))
    files = {name: filegen.png(width, 300)}  # a different width is a different photo
    body = start_batch(client, seeded, files)
    complete_all(client, body, files)
    drain()
    return body


def sheet_of(client, batch_id) -> dict:
    [s] = client.get(f"/api/v1/batches/{batch_id}/sheets").json()["data"]
    r = client.get(f"/api/v1/sheets/{s['id']}")
    assert r.status_code == 200, r.text
    return r.json()["data"]


def row(sheet, section, metric) -> dict:
    sec = next(s for s in sheet["sections"] if s["key"] == section)
    return next(r for r in sec["rows"] if r["metric"] == metric)


def patch(client, sheet, body) -> dict:
    r = client.patch(
        f"/api/v1/sheets/{sheet['id']}", json=body, headers={**idem(), "If-Match": f'"{sheet["version"]}"'}
    )
    assert r.status_code == 200, r.text
    return r.json()["data"]


def test_notebook_photo_to_database_sheet_files_and_email(client, seeded, monkeypatch, owner_engine):
    sign_in(client, seeded, "dev-reviewer")
    body = photo(client, seeded, monkeypatch, NOTEBOOK)

    # Read into the sheet of the date written on the page; supervisors and notes kept.
    s = sheet_of(client, body["batch_id"])
    assert s["report_date"] == "2026-10-02" and s["date_confirmed"] and s["state"] == "DRAFT"
    assert s["shifts"]["I"]["supervisor"] == "Ramesh Patel" and s["notes"][0]["label"] == "Beam fall m/c no."
    picks = row(s, "sulzer", "picks")
    assert [picks["cells"][sh]["value"] for sh in ("I", "II", "III")] == ["7026", "7087", "7071"]
    assert picks["total"] == "21184" and picks["cells"]["I"]["evidence"][0]["span_id"] == "p1-s6"
    eff = row(s, "sulzer", "total_eff_pct")  # calculated, exactly like the company's sheet
    assert eff["kind"] == "derived" and abs(float(eff["total"]) - 64.56) < 0.01
    with owner_engine.connect() as conn:  # stored in the database, one row per written value
        stored = conn.execute(
            select(t.shift_report_value).where(t.shift_report_value.c.report_id == uuid.UUID(s["id"]))
        ).all()
    assert len(stored) == 32  # 10 rows x 3 shifts + weft cut I and II; nothing calculated is stored
    assert {(v.section, v.metric, v.shift): str(v.value) for v in stored}[("sulzer", "picks", "II")] == "7087.0000"

    # Uncertain value (only two shift values written): approval is blocked until a person checks it.
    assert s["uncertain"] == 2 and not s["approvable"]
    r = client.post(f"/api/v1/sheets/{s['id']}/approve", headers={**idem(), "If-Match": f'"{s["version"]}"'})
    assert r.status_code == 422 and "Weft cut" in r.json()["error"]["fields"][0]["message"]
    s = patch(
        client,
        s,
        {
            "values": [{"section": "downtime", "metric": "weft_cut", "shift": "III", "value": "43.16"}],
            "confirm": [{"section": "downtime", "metric": "weft_cut", "shift": sh} for sh in ("I", "II")],
            "shifts": {
                "I": {"supervisor": "Ramesh Patel"},
                "II": {"supervisor": "Suresh"},
                "III": {"supervisor": "Mahesh"},
            },
        },
    )
    assert s["uncertain"] == 0 and s["approvable"] and row(s, "downtime", "weft_cut")["total"] == "114.06"
    r = client.post(f"/api/v1/sheets/{s['id']}/approve", headers={**idem(), "If-Match": f'"{s["version"]}"'})
    assert r.status_code == 200 and r.json()["data"]["state"] == "APPROVED"
    s = r.json()["data"]

    # A correction of an approved sheet needs a reason and is kept in the history.
    bad = client.patch(
        f"/api/v1/sheets/{s['id']}",
        headers={**idem(), "If-Match": f'"{s["version"]}"'},
        json={"values": [{"section": "sulzer", "metric": "picks", "shift": "I", "value": "7030"}]},
    )
    assert bad.status_code == 422
    s = patch(
        client,
        s,
        {
            "values": [{"section": "sulzer", "metric": "picks", "shift": "I", "value": "7030"}],
            "reason": "Supervisor corrected shift I picks",
        },
    )
    assert row(s, "sulzer", "picks")["total"] == "21188" and s["changes"][0]["old"] == "7026"

    # The next day: To date is the average of the daily totals of the month.
    day2 = [x.replace("2-Oct-26", "3-Oct-26") for x in NOTEBOOK]
    day2[day2.index("Picks 7026 7087 7071")] = "Picks 7000 7000 7000"
    s2 = sheet_of(client, photo(client, seeded, monkeypatch, day2, name="day-2.png")["batch_id"])
    assert s2["report_date"] == "2026-10-03" and row(s2, "sulzer", "picks")["to_date"] == "21094"  # (21188+21000)/2

    # The sheet list ("SQL sheet") with key figures and supervisors, filterable by supervisor.
    listed = client.get("/api/v1/sheets?supervisor=mahesh").json()["data"]
    assert [x["report_date"] for x in listed] == ["2026-10-02"] and listed[0]["figures"]["picks"] == "21188"

    # Downloads: Excel with formulas and a Data tab, CSV rows, printable PDF.
    x = client.get(f"/api/v1/sheets/{s['id']}/file?format=xlsx")
    assert x.status_code == 200 and x.headers["content-disposition"].endswith('Daily_Sheet_2026-10-02_Tapeline.xlsx"')
    wb = load_workbook(io.BytesIO(x.content))
    rep = wb["Report"]
    formulas = [c.value for r_ in rep.iter_rows() for c in r_ if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("SUM(D" in f for f in formulas) and any("*14.86" in f or "$" in f for f in formulas)
    assert wb["Data"].max_row > 20 and wb["Notes"]["B2"].value.startswith("B/F=")
    rows = list(
        csv.reader(io.StringIO(client.get(f"/api/v1/sheets/{s['id']}/file?format=csv").content.decode("utf-8-sig")))
    )
    picks_row = next(r_ for r_ in rows if r_[2] == "Sulzer production" and r_[3] == "Picks")
    assert picks_row[6:9] == ["7030", "7087", "7071"] and picks_row[10] == "21188"
    text = (
        PdfReader(io.BytesIO(client.get(f"/api/v1/sheets/{s['id']}/file?format=pdf").content)).pages[0].extract_text()
    )
    assert "Daily production sheet" in text and "21,188.00" in text and "Ramesh Patel" in text

    # Email: queued, sent by the worker with the PDF attached and the key figures in the message.
    mail = FakeEmailJs(monkeypatch)
    refused = client.post(
        f"/api/v1/sheets/{s['id']}/emails",
        headers=idem(),
        json={"to_email": "owner@example.com", "format": "pdf", "version": s["version"]},
    )
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "EMAIL_NOT_CONFIGURED"
    configure_email(client, seeded)
    for bad_address in ("not-an-email", "a@b"):
        r = client.post(
            f"/api/v1/sheets/{s['id']}/emails",
            headers=idem(),
            json={"to_email": bad_address, "format": "pdf", "version": s["version"]},
        )
        assert r.status_code == 422
    r = client.post(
        f"/api/v1/sheets/{s['id']}/emails",
        headers=idem(),
        json={"to_email": "owner@example.com", "format": "pdf", "version": s["version"]},
    )
    assert r.status_code == 202 and r.json()["data"]["state"] == "QUEUED", r.text
    drain()
    p = mail.sent[0]["template_params"]
    assert p["to_email"] == "owner@example.com" and "Picks: I 7,030.00" in p["message"] and "<table" in p["orders_html"]
    attached = base64.b64decode(p["pdf_file"].split(",", 1)[1])
    assert attached.startswith(b"%PDF") and p["attachment_name"] == "Daily_Sheet_2026-10-02_Tapeline.pdf"
    r = client.post(
        f"/api/v1/sheets/{s['id']}/emails",
        headers=idem(),
        json={"to_email": "owner@example.com", "format": "xlsx", "version": s["version"]},
    )
    drain()
    xp = mail.sent[1]["template_params"]
    assert xp["pdf_file"] == "" and xp["sheet_file"].startswith("data:application/vnd.openxmlformats")
    view = client.get(f"/api/v1/sheets/{s['id']}").json()["data"]
    assert [e["state"] for e in view["emails"]] == ["ACCEPTED", "ACCEPTED"]
    with owner_engine.connect() as conn:  # both fit EmailJS's 50 KB request limit (base64 adds a third)
        sizes = conn.execute(
            select(t.sheet_email.c.format, t.sheet_email.c.attachment_bytes).where(
                t.sheet_email.c.report_id == uuid.UUID(s["id"])
            )
        ).all()
    print("attachment sizes", sizes)
    assert all(b < 30_000 for _, b in sizes)


def test_second_page_fills_gaps_and_flags_differences(client, seeded, monkeypatch):
    sign_in(client, seeded, "dev-reviewer")
    first = photo(client, seeded, monkeypatch, NOTEBOOK[:9])  # date, supervisors, Sulzer rows
    s = sheet_of(client, first["batch_id"])
    second = [
        "2-Oct-26",
        "SULZER PROD",
        "Picks 7026 7090 7071",
        "No. of running looms 67.91 68.46 68.44",
        "Production in Kg 8919 8981 8955",
        "Avg width 3.33 3.33 3.33",
        "Downtime",
        "a) Mechanical 10 20 30",
    ]
    photo(client, seeded, monkeypatch, second, name="page-2.png")
    s2 = client.get(f"/api/v1/sheets/{s['id']}").json()["data"]  # the same sheet: one per department and day
    picks = row(s2, "sulzer", "picks")["cells"]
    assert picks["II"]["value"] == "7087" and picks["II"]["uncertain"] and "7090" in picks["II"]["note"]
    assert row(s2, "downtime", "mechanical")["total"] == "60" and len(s2["sources"]) == 2


def test_low_confidence_photo_and_missing_date_need_a_person(client, seeded, monkeypatch):
    sign_in(client, seeded, "dev-uploader")
    lines = [x for x in NOTEBOOK if not x.startswith("Date")]
    s = sheet_of(client, photo(client, seeded, monkeypatch, lines, confidence=0.6)["batch_id"])
    assert not s["date_confirmed"] and s["uncertain"] > 5 and not s["approvable"]
    s = patch(client, s, {"report_date": "2026-09-28"})  # uploaders may correct a draft
    assert s["report_date"] == "2026-09-28" and s["date_confirmed"]
    r = client.post(f"/api/v1/sheets/{s['id']}/approve", headers={**idem(), "If-Match": f'"{s["version"]}"'})
    assert r.status_code == 403  # only reviewers approve
    sign_in(client, seeded, "dev-viewer")  # Tapeline viewer: reads, cannot edit
    assert client.get(f"/api/v1/sheets/{s['id']}").status_code == 200
    assert client.patch(f"/api/v1/sheets/{s['id']}", json={}, headers={**idem(), "If-Match": '"1"'}).status_code == 403


def test_real_diary_page_mismatching_written_figures_highlight_the_misread_value(client, seeded, monkeypatch):
    """The real page writes Sulzer running looms as 80.04 (the target; shift I was 67.91): the written
    efficiencies do not fit it, so it is highlighted. Picks is confirmed by the matching total efficiency."""
    from tests.unit.test_daily_sheet import REAL_DIARY

    sign_in(client, seeded, "dev-reviewer")
    s = sheet_of(client, photo(client, seeded, monkeypatch, REAL_DIARY)["batch_id"])
    looms = row(s, "sulzer", "running_looms")["cells"]["I"]
    assert looms["value"] == "80.04" and looms["uncertain"] and "probably misread" in looms["note"]
    assert not row(s, "sulzer", "picks")["cells"]["I"]["uncertain"]
    assert not row(s, "normal_fibc", "running_looms")["cells"]["I"]["uncertain"]  # 53.41 fits its efficiencies
    assert any(n["label"] == "Check" and "403.81" in n["text"] for n in s["notes"])  # total that does not add up
    assert all(
        c["value"] is None
        for r_ in row(s, "sulzer", "picks")["cells"]
        for c in [row(s, "sulzer", "picks")["cells"][r_]]
        if r_ != "I"
    )  # nothing put in shifts II / III
    assert not s["approvable"]
