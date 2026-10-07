"""Pick reading register end to end, with the two real pages of 2 Oct 2026: photo -> OCR lines -> the day's register
in the database (one row per machine, time and shift) -> arithmetic checks -> review (correct / OK) -> approve ->
list -> Excel / PDF / CSV / SQL (loaded into SQLite to prove it runs) -> email with the file attached.
OCR is a stand-in returning lines in the Azure adapter's shape; EmailJS is the recorded transport of test_orders.
"""

import base64
import csv
import io
import sqlite3
import uuid

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select

from app.db import tables as t
from tests.conftest import idem, sign_in
from tests.integration.test_daily_sheets import photo
from tests.integration.test_orders import FakeEmailJs, NoteOcr, _infra, configure_email  # noqa: F401
from tests.unit.test_pick_register import PAGE_SHIFT_II, PAGE_SHIFT_III
from workers.runtime import run_one

pytestmark = [pytest.mark.db, pytest.mark.infra]
WORK = ["upload.scan", "upload.parse", "upload.extract", "register.email"]


def drain() -> None:
    while run_one(WORK, "test-worker"):
        pass


def register_of(client, batch_id) -> dict:
    items = client.get(f"/api/v1/batches/{batch_id}/sheets").json()["data"]
    [r] = [x for x in items if x["kind"] == "register"]
    view = client.get(f"/api/v1/registers/{r['id']}")
    assert view.status_code == 200, view.text
    return view.json()["data"]


def shift(reg, sh) -> dict:
    return next(s for s in reg["shifts"] if s["shift"] == sh)


def cell(reg, sh, machine, slot) -> dict:
    return next(m for m in shift(reg, sh)["machines"] if m["machine"] == machine)["cells"][slot]


def patch(client, reg, body) -> dict:
    r = client.patch(f"/api/v1/registers/{reg['id']}", json=body, headers={**idem(), "If-Match": f'"{reg["version"]}"'})
    assert r.status_code == 200, r.text
    return r.json()["data"]


def test_register_photos_to_database_checks_files_and_email(client, seeded, monkeypatch, owner_engine):
    sign_in(client, seeded, "dev-reviewer")
    first = photo(client, seeded, monkeypatch, PAGE_SHIFT_II, name="wgs02-shift2.png")
    photo(client, seeded, monkeypatch, PAGE_SHIFT_III, name="wgs02-shift3.png", width=421)
    reg = register_of(client, first["batch_id"])

    # Both pages went into the register of the written date; one per department and day.
    assert reg["register_date"] == "2026-10-02" and reg["date_confirmed"] and reg["state"] == "DRAFT"
    assert [s["shift"] for s in reg["sources"]] == ["II", "III"]
    c = cell(reg, "II", "27", 1)
    assert (c["reading"], c["picks"], c["source"]) == ("2090", "24", "read")
    assert cell(reg, "III", "26", 0)["status_label"] == "Beam fall"
    ii, iii = shift(reg, "II"), shift(reg, "III")
    assert [col["calculated"] for col in ii["columns"][1:]] == ["641", "682", "615", "625"]
    assert [col["match"] for col in iii["columns"][1:]] == [True, False, True, True]  # 04-00: 610 vs 590 written
    assert ii["total"] == "2563" and iii["total"] == "2442" and reg["day_total"] == "5005"
    assert iii["columns"][4]["stopped"] == 5 and iii["written_total"] == "687"
    with owner_engine.connect() as conn:  # one database row per written cell; nothing calculated is stored
        n = conn.execute(
            select(func.count()).select_from(t.pick_register_value)
            .where(t.pick_register_value.c.register_id == uuid.UUID(reg["id"]))
        ).scalar_one()  # fmt: skip
        stored = conn.execute(
            select(t.pick_register_value).where(
                t.pick_register_value.c.register_id == uuid.UUID(reg["id"]),
                t.pick_register_value.c.machine == "48",
                t.pick_register_value.c.shift == "II",
                t.pick_register_value.c.slot == 2,
            )
        ).one()
    assert n == 260 and (str(stored.reading), str(stored.picks)) == ("916.0000", "25.0000")

    # What a person must look at: 2 unclear numbers, 3 arithmetic checks. Approval is blocked until then.
    assert reg["uncertain"] == 2 and reg["checks"] == 4 and not reg["approvable"]  # incl. the 04-00 total
    assert "590" in iii["columns"][2]["check"]
    assert "1575 - 1357 (16-00) = 218" in cell(reg, "II", "39", 1)["checks"][0]["text"]
    assert "916 - 881 (18-00) = 35" in cell(reg, "II", "48", 2)["checks"][0]["text"]
    assert any("Machine 27, shift II" in f["text"] for f in reg["findings"])
    r = client.post(f"/api/v1/registers/{reg['id']}/approve", headers={**idem(), "If-Match": f'"{reg["version"]}"'})
    assert r.status_code == 422 and len(r.json()["error"]["fields"]) == 6

    # Review: the start reading of m/c 39 is corrected (the check clears by itself); the rest is right as written.
    reg = patch(client, reg, {"values": [{"shift": "II", "machine": "39", "slot": 0, "reading": "1557"}]})
    assert reg["checks"] == 3 and cell(reg, "II", "39", 1)["checks"] == []
    ok = [("II", "48", 2), ("III", "27", 2), ("III", "27", 3), ("II", "41", 3)]
    confirm = [{"shift": s, "machine": m, "slot": k} for s, m, k in ok] + [{"shift": "III", "slot": 2, "total": True}]
    reg = patch(client, reg, {"confirm": confirm})
    assert shift(reg, "III")["columns"][2]["accepted"] and shift(reg, "III")["columns"][2]["check"] is None
    assert reg["uncertain"] == 0 and reg["checks"] == 0 and reg["approvable"]
    assert cell(reg, "II", "48", 2)["accepted"]  # kept as accepted, visible on the cell
    r = client.post(f"/api/v1/registers/{reg['id']}/approve", headers={**idem(), "If-Match": f'"{reg["version"]}"'})
    assert r.status_code == 200 and r.json()["data"]["state"] == "APPROVED"
    reg = r.json()["data"]

    # A change to an approved register needs a reason and is kept in the history.
    bad = client.patch(
        f"/api/v1/registers/{reg['id']}",
        headers={**idem(), "If-Match": f'"{reg["version"]}"'},
        json={"values": [{"shift": "II", "machine": "48", "slot": 2, "picks": "35"}]},
    )
    assert bad.status_code == 422
    reg = patch(
        client, reg, {"values": [{"shift": "II", "machine": "48", "slot": 2, "picks": "35"}], "reason": "916-881 is 35"}
    )
    assert shift(reg, "II")["columns"][2]["calculated"] == "692" and reg["changes"][0]["old"] == "25"
    assert shift(reg, "II")["columns"][2]["match"] is False  # the worker's 682 used 25

    # The register list.
    listed = client.get("/api/v1/registers?date_from=2026-10-02&date_to=2026-10-02").json()["data"]
    assert [x["day_total"] for x in listed] == ["5015"] and listed[0]["present"] == ["II", "III"]

    # Downloads: Excel (formulas, Data and Checks tabs), CSV, PDF and SQL.
    x = client.get(f"/api/v1/registers/{reg['id']}/file?format=xlsx")
    assert x.status_code == 200 and x.headers["content-disposition"].endswith('Pick_Register_2026-10-02_Tapeline.xlsx"')
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Shift II", "Shift III", "Data", "Checks"]
    ws = wb["Shift III"]
    assert ws["A6"].value == "27" and ws["B6"].value == 2230 and ws["C6"].value == 2282 and ws["D6"].value == 24
    assert ws["K6"].value.startswith("=IF(COUNT(") and wb["Data"].max_row == 261
    rows = list(
        csv.reader(
            io.StringIO(client.get(f"/api/v1/registers/{reg['id']}/file?format=csv").content.decode("utf-8-sig"))
        )
    )
    assert ["2026-10-02", "Tapeline", "III", "02-00", "1", "27", "value", "2282", "24", "", "read"] in rows
    pdf = PdfReader(io.BytesIO(client.get(f"/api/v1/registers/{reg['id']}/file?format=pdf").content))
    text = " ".join(p.extract_text() for p in pdf.pages)
    assert "Hourly production reading register" in text and "2282 (24)" in text and len(pdf.pages) == 2
    sql = client.get(f"/api/v1/registers/{reg['id']}/file?format=sql")
    assert (
        sql.headers["content-type"].startswith("application/sql") and "attachment" in sql.headers["content-disposition"]
    )
    db = sqlite3.connect(":memory:")
    db.executescript(sql.text)
    db.executescript(sql.text)  # running it again replaces the day: no duplicates
    assert db.execute("SELECT COUNT(*) FROM pick_reading").fetchone()[0] == 260
    assert db.execute(
        "SELECT meter_reading, picks FROM pick_reading WHERE shift='II' AND machine='48' AND slot=2"
    ).fetchone() == (916, 35)
    assert db.execute("SELECT SUM(picks) FROM pick_reading WHERE shift='III'").fetchone()[0] == 2442
    assert db.execute("SELECT written_total FROM pick_reading_total WHERE shift='III' AND slot=2").fetchone()[0] == 590

    # Email: queued, sent by the worker with the file attached and the picks per shift in the message.
    mail = FakeEmailJs(monkeypatch)
    configure_email(client, seeded)
    r = client.post(
        f"/api/v1/registers/{reg['id']}/emails", headers=idem(),
        json={"to_email": "not-an-email", "format": "pdf", "version": reg["version"]},
    )  # fmt: skip
    assert r.status_code == 422
    for fmt in ("pdf", "sql"):
        r = client.post(
            f"/api/v1/registers/{reg['id']}/emails", headers=idem(),
            json={"to_email": "owner@example.com", "format": fmt, "version": reg["version"]},
        )  # fmt: skip
        assert r.status_code == 202 and r.json()["data"]["state"] == "QUEUED", r.text
        drain()
    p, q = mail.sent[0]["template_params"], mail.sent[1]["template_params"]
    assert "Shift III (00:00-08:00): picks 2442" in p["message"] and "<table" in p["orders_html"]
    assert base64.b64decode(p["pdf_file"].split(",", 1)[1]).startswith(b"%PDF")
    assert q["pdf_file"] == "" and q["sheet_file"].startswith("data:application/sql;base64,")
    assert b"INSERT INTO pick_reading" in base64.b64decode(q["sheet_file"].split(",", 1)[1])
    view = client.get(f"/api/v1/registers/{reg['id']}").json()["data"]
    assert [e["state"] for e in view["emails"]] == ["ACCEPTED", "ACCEPTED"]


def test_second_photo_of_a_page_confirms_and_flags_differences(client, seeded, monkeypatch):
    sign_in(client, seeded, "dev-reviewer")
    first = photo(client, seeded, monkeypatch, PAGE_SHIFT_II, name="a.png")
    again = [x.replace("2210 19", "2216 19") for x in PAGE_SHIFT_II]  # one number read differently
    photo(client, seeded, monkeypatch, again, name="b.png", width=421)
    reg = register_of(client, first["batch_id"])
    c = cell(reg, "II", "28", 2)
    assert c["reading"] == "2210" and c["uncertain"] and "2216" in c["note"]
    assert len(reg["sources"]) == 2 and cell(reg, "II", "29", 1)["evidence"][1]["upload_id"]  # confirmed twice


def test_access_scope_and_roles(client, seeded, monkeypatch):
    sign_in(client, seeded, "dev-uploader")
    reg = register_of(client, photo(client, seeded, monkeypatch, PAGE_SHIFT_III)["batch_id"])
    reg = patch(
        client, reg, {"values": [{"shift": "III", "machine": "56", "slot": 1, "reading": "100", "picks": "20"}]}
    )
    assert cell(reg, "III", "56", 1)["source"] == "manual"  # a machine row added by hand
    bad = client.patch(
        f"/api/v1/registers/{reg['id']}", headers={**idem(), "If-Match": f'"{reg["version"]}"'},
        json={"values": [{"shift": "III", "machine": "27", "slot": 0, "picks": "5"}]},
    )  # fmt: skip
    assert bad.status_code == 422  # the start reading has no picks
    r = client.post(f"/api/v1/registers/{reg['id']}/approve", headers={**idem(), "If-Match": f'"{reg["version"]}"'})
    assert r.status_code == 403  # only reviewers approve
    sign_in(client, seeded, "dev-viewer")  # reads, cannot edit
    assert client.get(f"/api/v1/registers/{reg['id']}").status_code == 200
    assert (
        client.patch(f"/api/v1/registers/{reg['id']}", json={}, headers={**idem(), "If-Match": '"1"'}).status_code
        == 403
    )


def test_photo_read_by_gemini_goes_into_the_register_and_sql(client, seeded, monkeypatch):
    """OCR_PROVIDER=gemini: the photo is transcribed by Gemini (a recorded answer here, no network) and lands in the
    register and its SQL file like any other page."""
    import json

    import httpx

    from app.ingestion.gemini import GeminiTranscriber

    def gemini(request: httpx.Request) -> httpx.Response:
        doc = {"legible": True, "lines": PAGE_SHIFT_III}
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": json.dumps(doc)}]}}]})

    sign_in(client, seeded, "dev-reviewer")
    reader = GeminiTranscriber(
        api_key="k", model="gemini-3.8-flash", http=httpx.Client(transport=httpx.MockTransport(gemini))
    )
    monkeypatch.setattr("workers.ingestion.get_ocr", lambda: reader)
    from tests import filegen
    from tests.integration.test_ingestion_pipeline import complete_all, start_batch

    files = {"register-photo.png": filegen.png(420, 300)}
    body = start_batch(client, seeded, files)
    complete_all(client, body, files)
    drain()
    reg = register_of(client, body["batch_id"])
    assert reg["sources"][0]["shift"] == "III" and shift(reg, "III")["total"] == "2442"
    sql = client.get(f"/api/v1/registers/{reg['id']}/file?format=sql").text
    db = sqlite3.connect(":memory:")
    db.executescript(sql)
    assert db.execute("SELECT COUNT(*) FROM pick_reading WHERE shift='III'").fetchone()[0] == 131


def test_the_same_photo_uploaded_again_is_not_read_twice(client, seeded, monkeypatch):
    """Workers re-send the same WhatsApp photo; a second reading (OCR / AI vary run to run) would only add
    differences to check. The exact same file is recorded and noted, not merged again."""
    sign_in(client, seeded, "dev-reviewer")
    first = photo(client, seeded, monkeypatch, PAGE_SHIFT_II, name="shift2.png")
    again = [x.replace("2210 19", "2216 19") for x in PAGE_SHIFT_II]  # a different reading of the same photo
    photo(client, seeded, monkeypatch, again, name="shift2-again.png")
    reg = register_of(client, first["batch_id"])
    assert cell(reg, "II", "28", 2)["reading"] == "2210" and not cell(reg, "II", "28", 2)["uncertain"]
    assert [s["values_read"] for s in reg["sources"]] == [129, 0]
    assert any(n["label"] == "Same photo" for n in reg["notes"])
