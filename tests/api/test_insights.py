"""Records list, dashboard, Excel export and control tower (FR11, FR12, FR14, A2, A9).

Covers TC13 (units separated), TC14 (zero target), TC23, TC24, TC25, TC29, TC57 (totals after correction)."""

import io
import uuid
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import insert, select, update
from sqlalchemy.exc import DBAPIError

from app.db import tables as t
from app.db.engine import tenant_tx
from app.domain.metrics import compute_metrics, format_pct
from tests.conftest import idem, sign_in
from workers.runtime import run_one

pytestmark = pytest.mark.db
DAY = "2026-09-27"  # the F1 fixture date


def add_record(
    owner_engine,
    seeded,
    *,
    day="2026-09-26",
    dept="TAPELINE",
    machine="T-01",
    qty="100",
    target="200",
    unit="m",
    status="RUNNING",
    stop=0,
    remarks="",
    archived=False,
    operator="Seeded",
):
    """Insert an approved record directly (owner connection) for aggregate scenarios."""
    rid, rev_id = uuid.uuid4(), uuid.uuid4()
    who = seeded.users["dev-reviewer"]
    with owner_engine.begin() as conn:
        conn.execute(
            insert(t.production_record).values(
                id=rid,
                tenant_id=seeded.tenant_id,
                department_id=seeded.departments[dept],
                current_revision_id=rev_id,
                production_date=date.fromisoformat(day),
                created_by=who,
                **(
                    {"state": "ARCHIVED", "archived_at": datetime.now(UTC), "archive_reason": "test archive"}
                    if archived
                    else {}
                ),
            )
        )
        conn.execute(
            insert(t.record_revision).values(
                id=rev_id,
                tenant_id=seeded.tenant_id,
                record_id=rid,
                number=1,
                production_date=date.fromisoformat(day),
                department_id=seeded.departments[dept],
                machine_id=seeded.machines[machine],
                operator_name=operator,
                production_qty=Decimal(qty),
                target_qty=Decimal(target),
                unit=unit,
                status=status,
                stop_minutes=stop,
                remarks=remarks,
                approval_state="APPROVED",
                created_by=who,
                approved_by=who,
                approved_at=datetime.now(UTC),
            )
        )
    return rid


def dash(client, **params):
    r = client.get("/api/v1/dashboard", params=params)
    assert r.status_code == 200, r.text
    return r.json()["data"]


# --- dashboard (FR12) ------------------------------------------------------------------------


def test_dashboard_reconciles_the_f1_fixture(client, seeded):  # TC24
    sign_in(client, seeded, "dev-reviewer")
    data = dash(client, date_from=DAY, date_to=DAY)
    assert data["metrics"] == [
        {
            "unit": "m",
            "production_qty": "4830.000",
            "target_qty": "6000.000",
            "achievement_pct": "80.5",
            "variance": "-1170.000",
            "record_count": 5,
        }
    ]
    assert data["record_count"] == 5 and data["stop_total_minutes"] == 75
    assert data["status_counts"] == {"RUNNING": 2, "COMPLETED": 1, "PENDING": 1, "HOLD": 1}
    assert data["status_shares"] == {"RUNNING": "40.0", "COMPLETED": "20.0", "PENDING": "20.0", "HOLD": "20.0"}
    by_dept = {x["department_name"]: x["achievement_pct"] for x in data["departments"]}
    assert by_dept == {
        "Tapeline": "83.3",
        "Warping": "98.0",
        "Lamination": "88.9",
        "Dispatch": "78.6",
        "Multifilament": "45.0",
    }
    assert data["data_version"] >= 1 and data["power_bi"]["state"] == "NOT_CONFIGURED"

    # The SQL aggregate and the domain formulas agree on the same rows (one formula set everywhere).
    rows = client.get("/api/v1/records", params={"date_from": DAY, "date_to": DAY}).json()["data"]
    domain = compute_metrics(
        [
            type(
                "R",
                (),
                {
                    "department_id": x["department"]["id"],
                    "unit": x["unit"],
                    "production_qty": Decimal(x["production_qty"]),
                    "target_qty": Decimal(x["target_qty"]),
                    "status": x["status"],
                    "stop_minutes": x["stop_minutes"],
                },
            )()
            for x in rows
        ]
    )
    assert format_pct(domain.by_unit[0].achievement_pct) == data["metrics"][0]["achievement_pct"]


def test_scope_follows_grants_and_filters_never_broaden_it(client, seeded):
    sign_in(client, seeded, "dev-viewer")  # Tapeline + Warping only
    data = dash(client, date_from=DAY, date_to=DAY)
    assert data["metrics"][0]["production_qty"] == "2230.000" and data["record_count"] == 2
    r = client.get(
        "/api/v1/dashboard", params={"date_from": DAY, "department_id": str(seeded.departments["LAMINATION"])}
    )
    assert r.status_code == 403
    assert client.get("/api/v1/records", params={"department_id": str(uuid.uuid4())}).status_code == 403


def test_units_are_never_added_together_and_zero_target_is_na(client, seeded, owner_engine):  # TC13, TC14
    add_record(owner_engine, seeded, day="2026-09-20", unit="kg", qty="150", target="0", machine="W-01", dept="WARPING")
    add_record(
        owner_engine, seeded, day="2026-09-20", unit="pcs", qty="12", target="10", machine="L-01", dept="LAMINATION"
    )
    add_record(owner_engine, seeded, day="2026-09-20", unit="m", qty="100", target="200")
    sign_in(client, seeded, "dev-reviewer")
    data = dash(client, date_from="2026-09-20", date_to="2026-09-20")
    assert [(x["unit"], x["production_qty"], x["achievement_pct"]) for x in data["metrics"]] == [
        ("m", "100.000", "50.0"),
        ("kg", "150.000", None),
        ("pcs", "12.000", "120.0"),
    ]


def test_archived_records_are_excluded_unless_asked_for(client, seeded, owner_engine):
    add_record(owner_engine, seeded, day="2026-09-21", qty="999", archived=True)
    sign_in(client, seeded, "dev-reviewer")
    assert dash(client, date_from="2026-09-21", date_to="2026-09-21")["record_count"] == 0
    assert dash(client, date_from="2026-09-21", date_to="2026-09-21", include_archived=True)["record_count"] == 1


def test_empty_scope_is_zero_and_na(client, seeded):  # TC25
    sign_in(client, seeded, "dev-reviewer")
    data = dash(client, date_from="2026-01-01", date_to="2026-01-31")
    assert data["record_count"] == 0 and data["metrics"] == [] and set(data["status_shares"].values()) == {None}


def test_a_correction_moves_the_totals(client, seeded):  # TC57 totals part
    sign_in(client, seeded, "dev-reviewer")
    rec_id = str(seeded.records[0])  # Tapeline 1250 m
    version = client.get(f"/api/v1/records/{rec_id}").json()["data"]["version"]
    r = client.post(
        f"/api/v1/records/{rec_id}/revisions",
        json={"fields": {"production_qty": "1300"}, "reason": "Recount at shift end"},
        headers={"If-Match": f'"{version}"', **idem()},
    )
    assert dash(client, date_from=DAY, date_to=DAY)["metrics"][0]["production_qty"] == "4830.000"  # still pending
    client.post(f"/api/v1/records/{rec_id}/revisions/{r.json()['data']['revision_id']}/approve", headers=idem())
    m = dash(client, date_from=DAY, date_to=DAY)["metrics"][0]
    assert (m["production_qty"], m["achievement_pct"]) == ("4880.000", "81.3")


@pytest.mark.parametrize(
    ("params", "code"),
    [
        ({"date_from": "2026-09-10", "date_to": "2026-09-01"}, "VALIDATION_FAILED"),
        ({"date_from": "2025-01-01", "date_to": "2026-09-01"}, "VALIDATION_FAILED"),
        ({"status": "STOPPED"}, "VALIDATION_FAILED"),
        ({"unit": "yd"}, "VALIDATION_FAILED"),
    ],
)
def test_invalid_filters_are_rejected(client, seeded, params, code):
    sign_in(client, seeded, "dev-reviewer")
    r = client.get("/api/v1/dashboard", params=params)
    assert r.status_code == 422 and r.json()["error"]["code"] == code


# --- records list (FR11, TC23) ---------------------------------------------------------------


def test_paging_is_stable_and_filters_combine(client, seeded, owner_engine):
    for i in range(30):
        add_record(owner_engine, seeded, day="2026-09-22", qty=str(100 + i), operator=f"Op {i:02d}")
    sign_in(client, seeded, "dev-viewer")
    seen, cursor = [], None
    while True:
        params = {"date_from": "2026-09-22", "date_to": "2026-09-22", "size": 25} | (
            {"cursor": cursor} if cursor else {}
        )
        page = client.get("/api/v1/records", params=params).json()
        seen += [x["id"] for x in page["data"]]
        cursor = page["next_cursor"]
        if not cursor:
            break
    assert len(seen) == len(set(seen)) == page["total"] == 30

    by_qty = client.get(
        "/api/v1/records",
        params={"date_from": "2026-09-22", "date_to": "2026-09-22", "unit": "m", "sort": "production_desc", "size": 25},
    ).json()["data"]
    assert by_qty[0]["production_qty"] == "129.000" and by_qty[0]["sync_state"] == "NOT_CONFIGURED"
    one = client.get(
        "/api/v1/records", params={"date_from": "2026-09-22", "date_to": "2026-09-22", "operator": "op 07"}
    ).json()
    assert one["total"] == 1 and one["data"][0]["operator_name"] == "Op 07"
    assert client.get("/api/v1/records", params={"sort": "production_desc"}).status_code == 400  # needs one unit
    assert client.get("/api/v1/records", params={"cursor": "garbage"}).status_code == 400
    assert client.get("/api/v1/records", params={"size": 30}).status_code == 400


def test_search_text_is_not_a_pattern(client, seeded, owner_engine):
    add_record(owner_engine, seeded, day="2026-09-23", remarks="50% done")
    add_record(owner_engine, seeded, day="2026-09-23", remarks="half done")
    sign_in(client, seeded, "dev-reviewer")
    r = client.get("/api/v1/records", params={"date_from": "2026-09-23", "date_to": "2026-09-23", "q": "%"}).json()
    assert [x["remarks"] for x in r["data"]] == ["50% done"]  # "%" is literal, not a wildcard


# --- Excel export (FR14, TC29) ----------------------------------------------------------------


@pytest.fixture
def storage_ready():
    from app.storage.objects import get_storage

    try:
        get_storage().ensure_bucket()
    except Exception as exc:  # noqa: BLE001
        pytest.skip(f"object storage unreachable: {type(exc).__name__}")


@pytest.mark.infra
def test_export_matches_the_snapshot_and_keeps_text_inert(client, seeded, owner_engine, storage_ready):
    hostile = ['=HYPERLINK("http://evil.invalid","x")', "+SUM(1,2)", "-2+3", "@cmd", "https://example.invalid"]
    for i, text in enumerate(hostile):
        add_record(owner_engine, seeded, day="2026-09-24", qty=f"{10 + i}.125", remarks=text)
    sign_in(client, seeded, "dev-reviewer")
    body = {"filter": {"date_from": "2026-09-24", "date_to": "2026-09-24"}}
    r = client.post("/api/v1/exports", json=body, headers=idem())
    assert r.status_code == 202 and r.json()["data"]["row_count"] == 5
    export_id = r.json()["data"]["export_id"]
    assert client.get(f"/api/v1/exports/{export_id}").json()["data"]["url"] is None  # not rendered yet
    while run_one(["export.render"], "test-worker"):
        pass
    data = client.get(f"/api/v1/exports/{export_id}").json()["data"]
    assert data["state"] == "READY" and data["row_count"] == 5 and data["url"]
    content = httpx.get(data["url"]).content
    import hashlib

    assert hashlib.sha256(content).hexdigest() == data["sha256"]

    wb = load_workbook(io.BytesIO(content))  # formulas would be kept as "=..." strings with data_type "f"
    ws = wb["Records"]
    rows = list(ws.iter_rows(min_row=2, values_only=False))
    assert len(rows) == 5
    remarks = {row[10].value for row in rows}
    assert remarks == set(hostile)
    assert all(row[10].data_type == "s" for row in rows)
    assert all(isinstance(row[4].value, float) for row in rows)  # quantities are numbers
    assert all(isinstance(row[0].value, datetime) and row[0].number_format == "yyyy-mm-dd" for row in rows)
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        assert all("<f>" not in z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/"))
    total = sum(Decimal(str(row[4].value)) for row in rows)
    summary = {c[0].value: c for c in wb["Summary"].iter_rows(min_row=2) if c[0].value}
    assert Decimal(str(summary["All selected"][2].value)) == total == Decimal("60.625")
    assert wb["Metadata"].protection.sheet


@pytest.mark.infra
def test_export_is_a_frozen_snapshot_and_scoped(client, seeded, owner_engine, storage_ready):
    sign_in(client, seeded, "dev-reviewer")
    r = client.post("/api/v1/exports", json={"filter": {"date_from": DAY, "date_to": DAY}}, headers=idem())
    export_id = uuid.UUID(r.json()["data"]["export_id"])
    with pytest.raises(DBAPIError), tenant_tx(seeded.tenant_id) as conn:  # snapshots are immutable
        conn.execute(update(t.export).where(t.export.c.id == export_id).values(row_count=1))
    with owner_engine.connect() as conn:
        snap = conn.execute(select(t.export.c.snapshot_json).where(t.export.c.id == export_id)).scalar_one()
    assert sorted(x["production_qty"] for x in snap) == ["1250.000", "1600.000", "450.000", "550.000", "980.000"]

    sign_in(client, seeded, "dev-viewer")
    assert client.post("/api/v1/exports", json={"filter": {}}, headers=idem()).status_code == 403
    assert client.get(f"/api/v1/exports/{export_id}").status_code == 403  # viewers have no export role
    sign_in(client, seeded, "dev-sender")
    assert client.get(f"/api/v1/exports/{export_id}").status_code == 200  # sender with all grants
    # Report snapshots exist since M6: an unknown report is simply not found.
    assert client.post("/api/v1/exports", json={"report_id": str(uuid.uuid4())}, headers=idem()).status_code == 404


# --- control tower (A2, A9) ------------------------------------------------------------------


def test_control_tower_shows_who_submitted(client, seeded, owner_engine):
    sign_in(client, seeded, "dev-reviewer")
    # 2026-09-27 is a Sunday (not a working day by default): F1 departments submitted, the rest not expected.
    sunday = client.get("/api/v1/control-tower", params={"date": DAY}).json()["data"]
    status = {x["code"]: x["status"] for x in sunday["departments"]}
    assert status == {
        "TAPELINE": "SUBMITTED",
        "WARPING": "SUBMITTED",
        "SULZER_FABRIC": "NOT_EXPECTED",
        "LAMINATION": "SUBMITTED",
        "MULTIFILAMENT": "SUBMITTED",
        "DISPATCH": "SUBMITTED",
        "PURCHASE": "NOT_EXPECTED",
    }
    assert sunday["totals"]["metrics"][0]["production_qty"] == "4830.000" and not sunday["working_day"]

    add_record(owner_engine, seeded, day="2026-09-25", dept="WARPING", machine="W-01")  # a Friday in the past
    friday = client.get("/api/v1/control-tower", params={"date": "2026-09-25"}).json()["data"]
    status = {x["code"]: x["status"] for x in friday["departments"]}
    assert status["WARPING"] == "SUBMITTED" and status["TAPELINE"] == "MISSING" and friday["past_cutoff"]
    assert friday["status_counts"]["MISSING"] == 6
    assert friday["integrations"]["google_sheets"]["state"] == "NOT_CONFIGURED"

    sign_in(client, seeded, "dev-viewer")  # scope follows grants
    assert {
        x["code"] for x in client.get("/api/v1/control-tower", params={"date": DAY}).json()["data"]["departments"]
    } == {"TAPELINE", "WARPING"}
